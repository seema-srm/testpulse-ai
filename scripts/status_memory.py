"""Build weekly synthetic project status reports and persist local report memory."""

from contextlib import closing
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import json
from pathlib import Path
import sqlite3
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill

from scripts.dashboard_data import ROOT, build_team_metrics

STATUS_MEMORY_DB = ROOT / "data" / "status_memory.sqlite3"


def _week_start(report_date: date) -> date:
    return report_date - timedelta(days=report_date.weekday())


def _carryover_sprints(data: dict[str, Any]) -> dict[str, set[str]]:
    carryovers: dict[str, set[str]] = {}
    for dependency in data["current"]["dependencies"].to_dict("records"):
        if dependency["dependency_type"] != "carryover":
            continue
        carryovers.setdefault(dependency["source_issue"], set()).add(
            dependency["observed_sprint"]
        )
    return carryovers


def build_weekly_status_report(
    data: dict[str, Any], report_date: date
) -> dict[str, Any]:
    """Create a deterministic weekly report from current synthetic data."""
    week_start = _week_start(report_date)
    team_metrics = build_team_metrics(data)
    current_issues = data["current"]["issues"]
    carryovers = _carryover_sprints(data)
    team_names = {team["id"]: team["name"] for team in data["teams"]}

    stuck_items = []
    blockers = []
    for issue in current_issues.to_dict("records"):
        observed_sprints = carryovers.get(issue["issue_key"], set())
        sprint_ids = observed_sprints | {issue["sprint_id"]}
        if len(sprint_ids) > 1:
            ordered_sprints = sorted(
                sprint_ids,
                key=lambda sprint_id: int(sprint_id.rsplit("-", 1)[-1]),
            )
            stuck_items.append(
                {
                    "issue_key": issue["issue_key"],
                    "team_id": issue["team_id"],
                    "team_name": team_names[issue["team_id"]],
                    "summary": issue["summary"],
                    "status": issue["status"],
                    "active": issue["status"] != "Done",
                    "priority": issue["priority"],
                    "stuck_sprints": len(ordered_sprints),
                    "sprint_ids": ordered_sprints,
                }
            )

        if issue["status"] != "Done" and str(
            issue.get("blocked", "")
        ).casefold() == "true":
            blockers.append(
                {
                    "issue_key": issue["issue_key"],
                    "team_id": issue["team_id"],
                    "team_name": team_names[issue["team_id"]],
                    "summary": issue["summary"],
                    "status": issue["status"],
                    "priority": issue["priority"],
                }
            )

    stuck_items.sort(key=lambda item: (-item["stuck_sprints"], item["issue_key"]))
    blockers.sort(key=lambda item: (item["team_name"], item["issue_key"]))

    teams = []
    for _, metrics in team_metrics.iterrows():
        team_id = str(metrics["id"])
        team_stuck_count = sum(
            item["team_id"] == team_id and item["active"] for item in stuck_items
        )
        risk_flags = []
        if int(metrics["critical_open_defects"]) > 0:
            risk_flags.append("Critical open defects")
        if team_stuck_count:
            risk_flags.append(f"{team_stuck_count} item(s) carried across sprints")
        if int(metrics["blocked"]) > 0:
            risk_flags.append("Blocked test executions")
        if int(metrics["failed"]) > 0:
            risk_flags.append("Failed test executions")

        risk_level = (
            "High"
            if int(metrics["critical_open_defects"]) > 0 or team_stuck_count > 0
            else "Medium"
            if int(metrics["open_defects"]) > 0
            or int(metrics["blocked"]) > 0
            or int(metrics["failed"]) > 0
            else "Low"
        )
        teams.append(
            {
                "team_id": team_id,
                "team_name": str(metrics["name"]),
                "risk_level": risk_level,
                "risk_flags": risk_flags,
                "completion_rate": float(metrics["completion_rate"]),
                "pass_rate": float(metrics["pass_rate"]),
                "open_defects": int(metrics["open_defects"]),
                "critical_open_defects": int(metrics["critical_open_defects"]),
                "failed_executions": int(metrics["failed"]),
                "blocked_executions": int(metrics["blocked"]),
            }
        )

    active_stuck_count = sum(item["active"] for item in stuck_items)
    return {
        "week_start": week_start.isoformat(),
        "week_end": (week_start + timedelta(days=6)).isoformat(),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "teams": teams,
        "blockers": blockers,
        "stuck_items": stuck_items,
        "summary": (
            f"{len(teams)} teams reviewed; {len(blockers)} active blocker(s); "
            f"{active_stuck_count} active item(s) carried across more than one sprint."
        ),
    }


def _connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS weekly_status_reports (
            week_start TEXT PRIMARY KEY,
            generated_at TEXT NOT NULL,
            report_json TEXT NOT NULL
        )
        """
    )
    return connection


def save_weekly_status_report(
    report: dict[str, Any], db_path: Path = STATUS_MEMORY_DB
) -> None:
    """Persist one report per week; refreshing the same week updates that entry."""
    with closing(_connect(db_path)) as connection:
        connection.execute(
            """
            INSERT INTO weekly_status_reports (week_start, generated_at, report_json)
            VALUES (?, ?, ?)
            ON CONFLICT(week_start) DO UPDATE SET
                generated_at = excluded.generated_at,
                report_json = excluded.report_json
            """,
            (
                report["week_start"],
                report["generated_at"],
                json.dumps(report, allow_nan=False),
            ),
        )
        connection.commit()


def get_latest_weekly_status_report(
    db_path: Path = STATUS_MEMORY_DB,
) -> dict[str, Any] | None:
    if not db_path.exists():
        return None
    with closing(_connect(db_path)) as connection:
        row = connection.execute(
            """
            SELECT report_json
            FROM weekly_status_reports
            ORDER BY week_start DESC
            LIMIT 1
            """
        ).fetchone()
    return json.loads(row["report_json"]) if row else None


def get_previous_weekly_status_report(
    week_start: str, db_path: Path = STATUS_MEMORY_DB
) -> dict[str, Any] | None:
    if not db_path.exists():
        return None
    with closing(_connect(db_path)) as connection:
        row = connection.execute(
            """
            SELECT report_json
            FROM weekly_status_reports
            WHERE week_start < ?
            ORDER BY week_start DESC
            LIMIT 1
            """,
            (week_start,),
        ).fetchone()
    return json.loads(row["report_json"]) if row else None


def format_weekly_status_markdown(
    report: dict[str, Any], previous_report: dict[str, Any] | None = None
) -> str:
    """Render a shareable weekly report with optional week-over-week changes."""
    previous_by_team = {
        team["team_id"]: team for team in (previous_report or {}).get("teams", [])
    }
    lines = [
        f"# Weekly Project Status - Week of {report['week_start']}",
        "",
        report["summary"],
        "",
        "## Team Status",
        "",
        "| Team | Risk | Completion | Pass Rate | Open Defects | Critical | Failed | Blocked |",
        "|---|---|---:|---:|---:|---:|---:|---:|",
    ]
    for team in report["teams"]:
        previous = previous_by_team.get(team["team_id"])
        pass_rate = f"{team['pass_rate']:.2f}%"
        lines.append(
            f"| {team['team_name']} | {team['risk_level']} | "
            f"{team['completion_rate']:.2f}% | {pass_rate} | "
            f"{team['open_defects']} | {team['critical_open_defects']} | "
            f"{team['failed_executions']} | {team['blocked_executions']} |"
        )

    if previous_report:
        lines.extend(["", "## Week-over-Week Changes", ""])
        previous_by_team = {
            team["team_id"]: team for team in previous_report["teams"]
        }
        for team in report["teams"]:
            previous = previous_by_team.get(team["team_id"])
            if previous is None:
                continue
            lines.append(
                f"- **{team['team_name']}**: pass rate "
                f"{team['pass_rate'] - previous['pass_rate']:+.2f} pts; "
                f"completion rate "
                f"{team['completion_rate'] - previous['completion_rate']:+.2f} pts; "
                f"open defects "
                f"{team['open_defects'] - previous['open_defects']:+d}; "
                f"critical defects "
                f"{team['critical_open_defects'] - previous['critical_open_defects']:+d}; "
                f"failed tests "
                f"{team['failed_executions'] - previous['failed_executions']:+d}; "
                f"blocked tests "
                f"{team['blocked_executions'] - previous['blocked_executions']:+d}."
            )

    lines.extend(["", "## Blockers", ""])
    if report["blockers"]:
        for blocker in report["blockers"]:
            lines.append(
                f"- **{blocker['issue_key']}** ({blocker['team_name']}, "
                f"{blocker['priority']}, {blocker['status']}): {blocker['summary']}"
            )
    else:
        lines.append("No active blockers are recorded.")

    lines.extend(["", "## Items Stuck Across More Than One Sprint", ""])
    if report["stuck_items"]:
        for item in report["stuck_items"]:
            lines.append(
                f"- **{item['issue_key']}** ({item['team_name']}) - "
                f"{item['stuck_sprints']} sprints; "
                f"{item['status']}{'' if item['active'] else ' (resolved)'}; "
                f"{item['priority']}: {item['summary']}"
            )
    else:
        lines.append("No multi-sprint carryover items are recorded.")
    lines.extend(
        [
            "",
            "_Risk levels are deterministic: High indicates a critical open defect "
            "or multi-sprint carryover; Medium indicates open defects, failed "
            "executions, or blocked executions; otherwise Low._",
        ]
    )
    return "\n".join(lines)


def week_over_week_frame(
    report: dict[str, Any], previous_report: dict[str, Any]
) -> pd.DataFrame:
    """Calculate team metric changes against the previous saved weekly report."""
    previous_by_team = {
        team["team_id"]: team for team in previous_report.get("teams", [])
    }
    rows = []
    for team in report["teams"]:
        previous = previous_by_team.get(team["team_id"])
        if previous is None:
            continue
        rows.append(
            {
                "Team": team["team_name"],
                "Pass Rate (pts)": round(
                    team["pass_rate"] - previous["pass_rate"], 2
                ),
                "Completion Rate (pts)": round(
                    team["completion_rate"] - previous["completion_rate"], 2
                ),
                "Open Defects": (
                    team["open_defects"] - previous["open_defects"]
                ),
                "Critical Defects": (
                    team["critical_open_defects"]
                    - previous["critical_open_defects"]
                ),
                "Failed Tests": (
                    team["failed_executions"] - previous["failed_executions"]
                ),
                "Blocked Tests": (
                    team["blocked_executions"] - previous["blocked_executions"]
                ),
            }
        )
    return pd.DataFrame(rows)


def team_status_frame(report: dict[str, Any]) -> pd.DataFrame:
    """Return a display-ready frame without changing the persisted report."""
    return pd.DataFrame(
        [
            {
                "Team": team["team_name"],
                "Risk": team["risk_level"],
                "Risk Flags": ", ".join(team["risk_flags"]) or "None",
                "Completion Rate": f"{team['completion_rate']:.2f}%",
                "Pass Rate": f"{team['pass_rate']:.2f}%",
                "Open Defects": team["open_defects"],
                "Critical": team["critical_open_defects"],
                "Failed Tests": team["failed_executions"],
                "Blocked Tests": team["blocked_executions"],
            }
            for team in report["teams"]
        ]
    )


def format_weekly_status_excel(
    report: dict[str, Any], previous_report: dict[str, Any] | None = None
) -> bytes:
    """Create a formatted Excel workbook for the weekly status report."""
    teams = team_status_frame(report)
    blockers = pd.DataFrame(
        [
            {
                "Team": item["team_name"],
                "Issue": item["issue_key"],
                "Title": item["summary"],
                "Priority": item["priority"],
                "Status": item["status"],
            }
            for item in report["blockers"]
        ],
        columns=["Team", "Issue", "Title", "Priority", "Status"],
    )
    stuck_items = pd.DataFrame(
        [
            {
                "Team": item["team_name"],
                "Issue": item["issue_key"],
                "Title": item["summary"],
                "Status": item["status"],
                "State": "Active" if item["active"] else "Resolved",
                "Priority": item["priority"],
                "Sprints": item["stuck_sprints"],
                "Sprint History": ", ".join(item["sprint_ids"]),
            }
            for item in report["stuck_items"]
        ],
        columns=[
            "Team",
            "Issue",
            "Title",
            "Status",
            "State",
            "Priority",
            "Sprints",
            "Sprint History",
        ],
    )
    sheets = {
        "Team Status": teams,
        "Blockers": blockers,
        "Multi-Sprint Items": stuck_items,
    }
    if previous_report:
        sheets["Week-over-Week"] = week_over_week_frame(report, previous_report)

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    fill_type="solid", fgColor="1E3A5F"
                )
            for column in worksheet.columns:
                column_letter = column[0].column_letter
                max_width = max(
                    (len(str(cell.value or "")) for cell in column),
                    default=10,
                )
                worksheet.column_dimensions[column_letter].width = min(
                    max(max_width + 2, 12), 48
                )
            if sheet_name == "Team Status":
                risk_fills = {
                    "High": "FECACA",
                    "Medium": "FED7AA",
                    "Low": "DCFCE7",
                }
                for row in range(2, worksheet.max_row + 1):
                    risk_cell = worksheet.cell(row, 2)
                    fill_color = risk_fills.get(risk_cell.value)
                    if fill_color:
                        risk_cell.fill = PatternFill(
                            fill_type="solid", fgColor=fill_color
                        )
            if sheet_name in {"Blockers", "Multi-Sprint Items"}:
                for row in range(2, worksheet.max_row + 1):
                    status_column = 5 if sheet_name == "Blockers" else 4
                    status_cell = worksheet.cell(row, status_column)
                    status_colors = {
                        "Blocked": ("FFEDD5", "9A3412"),
                        "Open": ("FFEDD5", "9A3412"),
                        "In Progress": ("DBEAFE", "1E40AF"),
                        "Done": ("DCFCE7", "166534"),
                        "PASS": ("DCFCE7", "166534"),
                        "FAIL": ("FEE2E2", "991B1B"),
                    }
                    if status_style := status_colors.get(status_cell.value):
                        fill_color, text_color = status_style
                        status_cell.fill = PatternFill(
                            fill_type="solid", fgColor=fill_color
                        )
                        status_cell.font = Font(color=text_color, bold=True)
    return output.getvalue()
