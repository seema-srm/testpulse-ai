from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from scripts.dashboard_data import load_dashboard_data
from scripts.status_memory import (
    build_weekly_status_report,
    format_weekly_status_excel,
    format_weekly_status_markdown,
    get_latest_weekly_status_report,
    get_previous_weekly_status_report,
    save_weekly_status_report,
    team_status_frame,
    week_over_week_frame,
)


def add_second_sprint_carryover(data):
    data["current"]["dependencies"] = pd.concat(
        [
            data["current"]["dependencies"],
            pd.DataFrame(
                [
                    {
                        "source_issue": "QAA-1",
                        "affected_project": "QAA",
                        "observed_sprint": "QAA-SPR-02",
                        "dependency_type": "carryover",
                    }
                ]
            ),
        ],
        ignore_index=True,
    )


def test_builds_weekly_report_and_finds_multi_sprint_carryover(project_files):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)
    add_second_sprint_carryover(data)

    report = build_weekly_status_report(data, date(2026, 7, 22))

    assert report["week_start"] == "2026-07-20"
    assert report["week_end"] == "2026-07-26"
    assert report["teams"][0]["team_name"] == "Team Alpha"
    assert report["teams"][0]["risk_level"] == "High"
    assert report["stuck_items"] == [
        {
            "issue_key": "QAA-1",
            "team_id": "alpha",
            "team_name": "Team Alpha",
            "summary": "Synthetic defect",
            "status": "Open",
            "active": True,
            "priority": "Critical",
            "stuck_sprints": 2,
            "sprint_ids": ["QAA-SPR-01", "QAA-SPR-02"],
        }
    ]


def test_local_memory_persists_reports_and_returns_previous_week(
    project_files, tmp_path
):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)
    db_path = tmp_path / "status-memory.sqlite3"

    first = build_weekly_status_report(data, date(2026, 7, 20))
    second = build_weekly_status_report(data, date(2026, 7, 27))
    save_weekly_status_report(first, db_path)
    save_weekly_status_report(second, db_path)

    assert get_latest_weekly_status_report(db_path)["week_start"] == "2026-07-27"
    assert get_previous_weekly_status_report("2026-07-27", db_path)["week_start"] == (
        "2026-07-20"
    )


def test_refreshing_a_week_updates_that_week_entry(project_files, tmp_path):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)
    db_path = tmp_path / "status-memory.sqlite3"
    report = build_weekly_status_report(data, date(2026, 7, 20))
    report["summary"] = "Updated report for the same week."

    save_weekly_status_report(report, db_path)
    save_weekly_status_report(report, db_path)

    assert get_latest_weekly_status_report(db_path)["summary"] == (
        "Updated report for the same week."
    )
    assert get_previous_weekly_status_report("2026-07-20", db_path) is None


def test_report_markdown_contains_team_risk_and_stuck_query(project_files):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)
    add_second_sprint_carryover(data)
    report = build_weekly_status_report(data, date(2026, 7, 20))

    markdown = format_weekly_status_markdown(report)
    frame = team_status_frame(report)

    assert "## Items Stuck Across More Than One Sprint" in markdown
    assert "QAA-1" in markdown
    assert "Team Alpha" in markdown
    assert "High" in frame.loc[0, "Risk"]


def test_weekly_reports_calculate_changes_for_all_status_metrics(project_files):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)
    previous = build_weekly_status_report(data, date(2026, 7, 20))
    current = build_weekly_status_report(data, date(2026, 7, 27))
    previous["teams"][0]["pass_rate"] = 70.0
    previous["teams"][0]["completion_rate"] = 30.0
    previous["teams"][0]["open_defects"] = 5
    previous["teams"][0]["critical_open_defects"] = 2
    previous["teams"][0]["failed_executions"] = 4
    previous["teams"][0]["blocked_executions"] = 1

    frame = week_over_week_frame(current, previous)
    markdown = format_weekly_status_markdown(current, previous)

    assert frame.loc[0, "Pass Rate (pts)"] == current["teams"][0]["pass_rate"] - 70.0
    assert frame.loc[0, "Open Defects"] == current["teams"][0]["open_defects"] - 5
    assert "## Week-over-Week Changes" in markdown
    assert "critical defects" in markdown


def test_excel_export_has_report_sheets_and_formatted_team_risk(project_files):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)
    previous = build_weekly_status_report(data, date(2026, 7, 20))
    current = build_weekly_status_report(data, date(2026, 7, 27))
    workbook_bytes = format_weekly_status_excel(current, previous)
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)

    assert workbook.sheetnames == [
        "Team Status",
        "Blockers",
        "Multi-Sprint Items",
        "Week-over-Week",
    ]
    team_status = workbook["Team Status"]
    assert team_status["A1"].value == "Team"
    assert team_status["B1"].value == "Risk"
    assert team_status["A2"].value == "Team Alpha"
    assert team_status.freeze_panes == "A2"
    assert team_status.auto_filter.ref is not None
    assert team_status["B2"].fill.fgColor.rgb.endswith("FECACA")


def test_non_stuck_issue_is_not_reported_as_multi_sprint(project_files):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)

    report = build_weekly_status_report(data, date(2026, 7, 20))

    assert report["stuck_items"] == []


def test_completed_multi_sprint_item_remains_in_history(project_files):
    history, teams_config = project_files
    data = load_dashboard_data(history, history.parent / "no-events.jsonl", teams_config)
    add_second_sprint_carryover(data)
    data["current"]["issues"].loc[
        data["current"]["issues"]["issue_key"] == "QAA-1", "status"
    ] = "Done"

    report = build_weekly_status_report(data, date(2026, 7, 20))

    assert report["stuck_items"][0]["issue_key"] == "QAA-1"
    assert report["stuck_items"][0]["active"] is False
